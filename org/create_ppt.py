from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_title_slide(title, subtitle):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5)
    )
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(4), Inches(12.333), Inches(1)
    )
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = subtitle
    subtitle_para.font.size = Pt(24)
    subtitle_para.alignment = PP_ALIGN.CENTER


def add_content_slide(title, bullets):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(12.333), Inches(1)
    )
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(36)
    title_para.font.bold = True

    content_box = slide.shapes.add_textbox(
        Inches(0.7), Inches(1.3), Inches(12), Inches(5.5)
    )
    content_frame = content_box.text_frame
    content_frame.word_wrap = True

    for bullet in bullets:
        p = (
            content_frame.paragraphs[0]
            if not content_frame.paragraphs
            else content_frame.add_paragraph()
        )
        p.text = bullet
        p.font.size = Pt(22)
        p.space_after = Pt(12)


add_title_slide("Low-Light Image Enhancement", "Using Pixel-Adaptive Gamma Correction")

add_content_slide(
    "Overview",
    [
        "This project implements a low-light image enhancement technique",
        "Published in Pattern Recognition journal (2023)",
        "Key idea: Transform low-light problem into image dehazing problem",
        "Uses gamma correction prior in mixed color spaces",
        "Achieves better results by adapting to each pixel individually",
    ],
)

add_content_slide(
    "Key Concepts",
    [
        "Dark Channel Prior: Used to estimate atmospheric light in dehazing",
        "Gamma Correction: Non-linear transformation to adjust brightness",
        "Transmission Map: Estimates how much light reaches the camera",
        "Retinex Theory: Models human visual system for color perception",
        "Normalization: Scales pixel values to consistent range",
    ],
)

add_content_slide(
    "Step 1: Image Preprocessing",
    [
        "Input image is read from disk in color format",
        "Pixel values converted from 0-255 to 0-1 floating point",
        "Gaussian blur applied to create a denoised version",
        "This denoised image helps in accurate light estimation",
        "Process: i2f() converts integer to float, f2i() converts back",
    ],
)

add_content_slide(
    "Step 2: Computing Atmospheric Light",
    [
        "Uses Dark Channel Prior method from dehazing research",
        "Finds the darkest pixel in each local neighborhood",
        "Selects top 0.1% brightest pixels from dark channel",
        "Calculates median RGB values of these bright pixels",
        "This represents the atmospheric light 'A' in the scene",
    ],
)

add_content_slide(
    "Step 3: Image Normalization",
    [
        "Divide each color channel by atmospheric light value",
        "This normalizes the image based on lighting conditions",
        "After normalization, bright areas become more prominent",
        "Normalization function scales each channel independently",
        "Ensures all values stay within 0 to 1 range",
    ],
)

add_content_slide(
    "Step 4: Pixel-Adaptive Gamma Correction",
    [
        "Gamma correction adjusts image brightness non-linearly",
        "Traditional gamma uses single value for entire image",
        "This method adapts gamma value for each pixel",
        "Uses exponential function based on maximum pixel intensity",
        "Gamma ranges from 1 (no change) to gamma_max (maximum adjustment)",
    ],
)

add_content_slide(
    "Step 5: Estimating Transmission Map",
    [
        "Transmission map indicates how much light passes through",
        "Higher transmission means less atmospheric interference",
        "Calculated using intensity and maximum color values",
        "Before and after gamma correction are compared",
        "Minimum transmission limited to 0.1 to prevent artifacts",
    ],
)

add_content_slide(
    "Step 6: Recovering Enhanced Image",
    [
        "Uses formula similar to image dehazing",
        "Reverses the atmospheric scattering model",
        "Combines original image, transmission map, and atmospheric light",
        "Recover function applies the recovery equation",
        "Output clipped to valid range 0-1",
    ],
)

add_content_slide(
    "Step 7: Final Adjustment",
    [
        "Percentile-based adjustment for better visual quality",
        "Uses 99.5th percentile as high value",
        "Uses 0.5th percentile as low value",
        "Stretches the contrast across the full range",
        "Final result saved as output image",
    ],
)

add_content_slide(
    "Workflow Summary",
    [
        "1. Load and preprocess image (float conversion)",
        "2. Compute atmospheric light using dark channel",
        "3. Normalize image by dividing with atmospheric light",
        "4. Apply pixel-adaptive gamma correction",
        "5. Estimate transmission map from corrected image",
        "6. Recover enhanced image using transmission",
        "7. Apply final percentile adjustment",
        "8. Save the enhanced output image",
    ],
)

add_content_slide(
    "Why This Works",
    [
        "Low-light images have similar characteristics to hazy images",
        "By inverting the problem, we can use dehazing techniques",
        "Gamma correction brightens dark areas effectively",
        "Pixel-adaptive approach handles varying light conditions",
        "Combines multiple color spaces (RGB, HSV, HSI concepts)",
    ],
)

add_content_slide(
    "Parameters",
    [
        "gamma_max: Controls maximum gamma correction (default: 6.0)",
        "Higher values brighten darker images more aggressively",
        "erosion_window: Size for dark channel calculation (15x15)",
        "n_bins: Histogram bins for threshold calculation (200)",
        "T_min: Minimum transmission value (0.1)",
    ],
)

add_content_slide(
    "Usage",
    [
        "Run from command line with input and output directories",
        "Optional gamma_max parameter for tuning brightness",
        "Processes all common image formats (jpg, png, bmp, tiff, webp)",
        "Saves enhanced images as PNG files",
        "Reports processing time for each image",
    ],
)

add_title_slide("Thank You", "Questions?")

prs.save("Low_Light_Image_Enhancement_Presentation.pptx")
print("PowerPoint presentation saved successfully!")
